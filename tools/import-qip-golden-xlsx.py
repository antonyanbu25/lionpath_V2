#!/usr/bin/env python3
"""Dump scoring-sheet xlsx rows as GoldenExample JSON (stdout).

Usage:
  python3 tools/import-qip-golden-xlsx.py "/path/to/sheet.xlsx" \\
      --motion new_business --source nb_eu_scoring_sheet

Paste the printed objects into worker/src/postcall/golden-set/examples.ts.

Expected sheet shape (one call per tab):
  A2 = SE name, B2 = call type label, A3 or A4 = Engage/recording URL
  C = category, D = theme label, E = theme key, F = credit, G = score (0/1/2)
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

CALL_TYPE_MAP = {
    "demo": "demo",
    "demo - plan upgrade": "demo_plan_upgrade",
    "plan upgrade": "demo_plan_upgrade",
    "ai vision demo": "ai_vision_demo",
    "discovery": "discovery",
}


def slug(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")
    return s or "call"


def infer_call_type(raw: object) -> str:
    label = str(raw or "").strip().lower()
    return CALL_TYPE_MAP.get(label, re.sub(r"[^a-z0-9]+", "_", label).strip("_") or "demo")


def looks_like_url(val: object) -> bool:
    return isinstance(val, str) and val.startswith("http")


def dump_workbook(path: Path, motion: str, source: str) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    out: list[dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        se_name = str(ws["A2"].value or sheet_name).strip()
        call_type = infer_call_type(ws["B2"].value)
        recording = ""
        for coord in ("A3", "A4", "A5"):
            val = ws[coord].value
            if looks_like_url(val):
                recording = str(val).split("?")[0]
                break
        theme_bands: dict[str, int] = {}
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7):
            key_cell, score_cell = row[4], row[6]
            if key_cell.value is None or score_cell.value is None:
                continue
            key = str(key_cell.value).strip()
            try:
                band = int(round(float(score_cell.value)))
            except (TypeError, ValueError):
                continue
            if band not in (0, 1, 2) or not key:
                continue
            theme_bands[key] = band
        if not theme_bands:
            continue
        out.append(
            {
                "id": f"{slug(se_name)}-{motion[:3]}-{call_type.replace('_', '-')}",
                "seName": se_name,
                "seMotion": motion,
                "callType": call_type,
                "recordingUrl": recording,
                "sourceFile": source,
                "scoredAt": "2026-09",
                "themeBands": theme_bands,
            }
        )
    return out


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--motion", choices=("new_business", "expansion"), required=True)
    p.add_argument("--source", default="imported_scoring_sheet")
    args = p.parse_args()
    path = Path(args.xlsx)
    if not path.exists():
        print(f"missing {path}", file=sys.stderr)
        return 1
    examples = dump_workbook(path, args.motion, args.source)
    print(json.dumps(examples, indent=2))
    print(
        f"\n# {len(examples)} call(s). Append to GOLDEN_EXAMPLES in worker/src/postcall/golden-set/examples.ts",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
